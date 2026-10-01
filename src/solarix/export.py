"""Exporta a tabela de fechamento (entregável 2) em xlsx: Resumo, Fechamento e Comissão por baixa."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .compare import COMPONENTES
from .config import TOLERANCIA_RS

ROTULOS = {"comissao": "Comissão s/ vendas", "bonus_meta": "Bônus de meta", "rampagem": "Compl. rampagem",
           "override": "Override", "total": "Total"}

# Leitura de cada divergência já conhecida (ids de docs/decisoes.md e docs/perguntas_carla.md)
OBSERVACOES = {
    ("Érica Lins", "2026-09"): "Rampagem não consta no controle: 2º mês de casa, garantia de R$ 1.500 menos comissão de "
                               "R$ 190,16 (ponto I1)",
    ("Bruno Sato", "2026-09"): "Comissão não explicada; pedir a conta por venda (pergunta Q4)",
    ("Camila Duarte", "2026-08"): "Comissão não explicada; valor redondo sugere ajuste manual (pergunta Q5)",
}

NUM = "#,##0.00"
TOLERANCIA_BR = f"{TOLERANCIA_RS:.2f}".replace(".", ",")  # texto em formato brasileiro
CINZA = PatternFill("solid", fgColor="D9D9D9")
VERMELHO = PatternFill("solid", fgColor="F8CBAD")
VERDE = PatternFill("solid", fgColor="E2EFDA")
FINA = Side(style="thin", color="999999")


def _cabecalho(ws, linha: int, textos: list[str]):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(bold=True)
        c.fill = CINZA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=FINA)


def _aba_resumo(ws, comp: pd.DataFrame):
    """Total do período por comissionado, nosso x Carla, e a lista das linhas que não batem."""
    ws.append(["Fechamento de comissões Solarix: agosto e setembro/2026"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Valores do cálculo de referência em Python, validado contra a RevTrack (diferenças de centavos por "
               f"arredondamento). Uma linha bate quando todos os componentes ficam a até R$ {TOLERANCIA_BR} do "
               f"controle da Carla."])
    ws.append([f"Linhas que batem: {int(comp['bateu'].sum())} de {len(comp)}."])
    ws.append([])
    _cabecalho(ws, 5, ["Comissionado", "Total ago + set (nosso)", "Total ago + set (Carla)", "Diferença", "Linhas que não batem"])
    por_pessoa = comp.groupby("nome")[["total_nosso", "total_carla"]].sum().round(2)
    nao_bate = comp[~comp["bateu"]].groupby("nome")["competencia"].apply(", ".join)
    for nome, r in por_pessoa.iterrows():
        dif = round(r["total_nosso"] - r["total_carla"], 2)
        ws.append([nome, r["total_nosso"], r["total_carla"], dif, nao_bate.get(nome, "")])
        if nome in nao_bate.index:
            for c in ws[ws.max_row]:
                c.fill = VERMELHO
    t = por_pessoa.sum().round(2)
    ws.append(["Total", t["total_nosso"], t["total_carla"], round(t["total_nosso"] - t["total_carla"], 2), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = VERDE
    for row in ws.iter_rows(min_row=6, min_col=2, max_col=4):
        for c in row:
            c.number_format = NUM
    for letra, larg in zip("ABCDE", (24, 24, 24, 14, 24), strict=True):
        ws.column_dimensions[letra].width = larg
    ws.freeze_panes = "A6"


def _aba_fechamento(ws, comp: pd.DataFrame):
    """Uma linha por comissionado x mês: nosso, Carla e diferença de cada componente, mais totais."""
    # faixa superior com o nome de cada componente
    ws.cell(row=1, column=1, value="Comissionado")
    ws.cell(row=1, column=2, value="Competência")
    col = 3
    for k in COMPONENTES:
        ws.cell(row=1, column=col, value=ROTULOS[k])
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        col += 3
    ws.cell(row=1, column=col, value="Bateu?")
    ws.cell(row=1, column=col + 1, value="Observação")
    sub = ["", ""] + ["Nosso", "Carla", "Dif."] * len(COMPONENTES) + ["", ""]
    for i in range(1, len(sub) + 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True), CINZA
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, t in enumerate(sub, start=1):
        c = ws.cell(row=2, column=i, value=t)
        c.font, c.fill = Font(bold=True), CINZA
        c.alignment = Alignment(horizontal="center")

    dados = comp.sort_values(["nome", "competencia"])
    cols_num = [f"{k}_{s}" for k in COMPONENTES for s in ("nosso", "carla", "dif")]
    for r in dados.itertuples(index=False):
        d = r._asdict()
        obs = OBSERVACOES.get((d["nome"], d["competencia"]), "" if d["bateu"] else "Ver colunas com diferença")
        ws.append([d["nome"], d["competencia"], *[d[c] for c in cols_num], "Sim" if d["bateu"] else "NÃO", obs])
        if not d["bateu"]:
            for c in ws[ws.max_row]:
                c.fill = VERMELHO
    ultima = ws.max_row
    # linha de totais (soma das 16 linhas)
    ws.append(["Total", ""] + [round(float(dados[c].sum()), 2) for c in cols_num] + ["", ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = VERDE
        c.border = Border(top=FINA)
    for row in ws.iter_rows(min_row=3, max_row=ultima + 1, min_col=3, max_col=2 + len(cols_num)):
        for c in row:
            c.number_format = NUM
    larguras = {1: 24, 2: 13, 3 + len(cols_num): 8, 4 + len(cols_num): 62}
    for i in range(1, 5 + len(cols_num)):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(i, 11)
    ws.freeze_panes = "C3"


def exportar_fechamento(comp: pd.DataFrame, por_baixa: pd.DataFrame, caminho):
    """Grava o xlsx final: abas Resumo, Fechamento (por comissionado e mês) e Comissão por baixa."""
    wb = Workbook()
    _aba_resumo(wb.active, comp)
    wb.active.title = "Resumo"
    _aba_fechamento(wb.create_sheet("Fechamento"), comp)
    ws = wb.create_sheet("Comissão por baixa")
    _cabecalho(ws, 1, ["Baixa (ERP)", "Venda", "Competência", "Comissão (antes do split)"])
    for r in por_baixa.sort_values(["competencia", "n_baixa"]).itertuples(index=False):
        ws.append([int(r.n_baixa), int(r.id_venda), r.competencia, float(r.comissao)])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = NUM
    for letra, larg in zip("ABCD", (14, 10, 14, 26), strict=True):
        ws.column_dimensions[letra].width = larg
    ws.freeze_panes = "A2"
    wb.save(caminho)
